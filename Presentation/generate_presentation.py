"""Generate the editable ten-slide deck, equation artwork and presenter guide."""
from pathlib import Path
import hashlib
import json
import math

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from xml.sax.saxutils import escape

from speaker_content import NOTES, FEATURE_COVERAGE


HERE = Path(__file__).resolve().parent
ASSETS = HERE / "assets"
ROOT = HERE.parent
BG = "10181E"
PANEL = "19262D"
WHITE = "F0E9DA"
MUTED = "ACB9BA"
CYAN = "58CFD1"
GOLD = "D5A45E"
RED = "F39472"
EDGE = "42545B"
FONT = "Arial"
DISPLAY = "Bahnschrift"
WIDTH, HEIGHT = 1600, 900


def rgb(value):
    return RGBColor.from_string(value)


def inch(value):
    return Inches(value / 100)


def make_assets():
    crops = {
        "aircraft": (340, 190, 990, 535),
        "aircraft_wire": (340, 190, 990, 535),
        "harvester": (330, 160, 1020, 560),
        "tank": (345, 150, 955, 570),
        "rocks": (280, 70, 1060, 635),
        "base": (190, 100, 1100, 640),
        "worker": (465, 155, 805, 585),
        "beacon": (300, 90, 1040, 630),
        "winch": (340, 135, 1030, 520),
        "mouth": (290, 30, 1020, 680),
        "telemetry": (18, 155, 247, 285),
        "radar": (1095, 140, 1267, 355),
        "status": (18, 17, 1260, 132),
        "action": (320, 609, 960, 661),
        "shadow_detail": (430, 275, 1060, 660),
        "dust_detail": (240, 175, 1080, 615),
    }
    for name, box in crops.items():
        origin = name if name in {"aircraft", "aircraft_wire", "harvester", "tank", "rocks", "base", "worker", "beacon", "winch", "mouth"} else "scene"
        if name == "shadow_detail": origin = "controls"
        if name == "dust_detail": origin = "breach"
        if name == "radar": origin = "pursuit"
        with Image.open(ASSETS / f"{origin}.png") as image:
            image.crop(box).save(ASSETS / f"{name}_crop.png")


class Deck:
    def __init__(self):
        self.presentation = Presentation()
        self.presentation.slide_width = Inches(16)
        self.presentation.slide_height = Inches(9)
        self.presentation.core_properties.title = "Arrakis: Harvester Down"
        self.presentation.core_properties.subject = "CSE 4102 - Computer Graphics and Image Processing Laboratory"
        self.presentation.core_properties.author = "Asique Ehetasamul Haque"
        self.presentation.core_properties.keywords = "Procedural graphics, controllable rescue, OpenGL, CSE 4102"
        self.slides = []
        self.manifest = []
        self.fonts = {}

    def new(self, title, subtitle="", section="", chrome=True):
        slide = self.presentation.slides.add_slide(self.presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(BG)
        self.slide = slide
        self.record = {"number": len(self.slides) + 1, "title": title, "texts": [], "images": [], "boxes": []}
        self.manifest.append(self.record)
        self.slides.append(slide)
        if chrome:
            self.text(60, 35, 1370, 24, section.upper(), 12, GOLD, True)
            self.text(60, 76, 1460, 65, title, 36, WHITE, True, DISPLAY)
            if subtitle: self.text(62, 144, 1470, 36, subtitle, 16, MUTED)
            self.line(60, 850, 1540, 850, EDGE, 1)
            self.text(60, 864, 1240, 24, "ARRAKIS / HARVESTER DOWN     |     ASIQUE EHETASAMUL HAQUE     |     CSE 4102", 10, MUTED)
            self.text(1440, 860, 100, 30, f"{len(self.slides):02d} / 10", 13, GOLD, True, align=PP_ALIGN.RIGHT)
        return slide

    def bounds(self, x, y, w, h, kind):
        if x < -0.1 or y < -0.1 or x + w > WIDTH + 0.1 or y + h > HEIGHT + 0.1:
            raise ValueError(f"Slide {len(self.slides)} out-of-bounds {kind}: {(x, y, w, h)}")
        self.record["boxes"].append({"kind": kind, "x": x, "y": y, "width": w, "height": h})

    def rect(self, x, y, w, h, fill=PANEL, line=None, radius=False):
        self.bounds(x, y, w, h, "shape")
        shape = self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, inch(x), inch(y), inch(w), inch(h))
        shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)
        if line: shape.line.color.rgb = rgb(line); shape.line.width = Pt(1)
        else: shape.line.fill.background()
        if radius: shape.adjustments[0] = 0.05
        return shape

    def line(self, x1, y1, x2, y2, colour=EDGE, width=1.5):
        shape = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, inch(x1), inch(y1), inch(x2), inch(y2))
        shape.line.color.rgb = rgb(colour); shape.line.width = Pt(width)
        return shape

    def text(self, x, y, w, h, text, size=18, colour=WHITE, bold=False, font=FONT, align=PP_ALIGN.LEFT):
        self.bounds(x, y, w, h, "text")
        shape = self.slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h))
        frame = shape.text_frame
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.word_wrap = True; frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.vertical_anchor = MSO_ANCHOR.TOP
        for index, row in enumerate(text.split("\n")):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = row
            paragraph.alignment = align
            paragraph.line_spacing = 1.0
            paragraph.space_before = Pt(0); paragraph.space_after = Pt(0)
            for run in paragraph.runs:
                run.font.name = font; run.font.size = Pt(size); run.font.bold = bold
                run.font.color.rgb = rgb(colour)
        self.record["texts"].append(text)
        return shape

    def image(self, name, x, y, w, h, fill=True):
        self.bounds(x, y, w, h, "image")
        path = ASSETS / (name if "." in name else f"{name}.png")
        with Image.open(path) as image: iw, ih = image.size
        if not fill:
            scale = min(w / iw, h / ih)
            x += (w - iw * scale) / 2; y += (h - ih * scale) / 2
            w, h = iw * scale, ih * scale
        shape = self.slide.shapes.add_picture(str(path), inch(x), inch(y), inch(w), inch(h))
        if fill:
            source_ratio, target_ratio = iw / ih, w / h
            if source_ratio > target_ratio:
                shape.crop_left = shape.crop_right = (1 - target_ratio / source_ratio) / 2
            else:
                shape.crop_top = shape.crop_bottom = (1 - source_ratio / target_ratio) / 2
        self.record["images"].append(path.name)
        return shape

    def equation(self, name, expression, x, y, w, h, size=22, colour=WHITE):
        path = ASSETS / f"eq_{name}.png"
        fig = plt.figure(figsize=(14, 1.2), dpi=180)
        fig.patch.set_alpha(0)
        fig.text(0.02, 0.30, f"${expression}$", fontsize=size, color=f"#{colour}")
        fig.savefig(path, dpi=180, transparent=True, bbox_inches="tight", pad_inches=0.025)
        plt.close(fig)
        return self.image(path.name, x, y, w, h, False)

    def arrow(self, x, y, w=26, colour=CYAN):
        shape = self.slide.shapes.add_shape(MSO_SHAPE.CHEVRON, inch(x), inch(y), inch(w), inch(30))
        shape.fill.solid(); shape.fill.fore_color.rgb = rgb(colour); shape.line.fill.background()

    def cue(self, text, x=60, y=798, w=1480):
        self.rect(x, y, w, 40, PANEL)
        self.rect(x, y, 4, 40, CYAN)
        self.text(x + 16, y + 8, w - 25, 26, text, 15, CYAN)


def create_deck():
    make_assets()
    d = Deck()

    # 1. Required introduction with the complete student identity.
    d.new("Introduction", chrome=False)
    d.image("breach", 760, 0, 840, 900)
    d.rect(0, 0, 805, 900, BG)
    d.rect(802, 0, 5, 900, GOLD)
    d.text(60, 45, 680, 35, "01 / INTRODUCTION", 13, GOLD, True)
    d.text(60, 103, 710, 140, "ARRAKIS", 78, WHITE, True, DISPLAY)
    d.text(65, 230, 710, 60, "HARVESTER DOWN", 33, GOLD, True, DISPLAY)
    d.text(65, 315, 680, 89, "Fly an ornithopter. Rescue the crew.\nEscape a pursuing sandworm.", 23)
    d.line(65, 423, 740, 423, EDGE)
    d.text(65, 449, 695, 50, "Asique Ehetasamul Haque", 25, WHITE, True)
    d.text(65, 505, 700, 40, "Roll: 2107096  |  Year: 4th  |  Semester: 1st  |  Group: B2", 16, MUTED)
    d.text(65, 563, 710, 78, "CSE - 4102\nComputer Graphics and Image Processing Laboratory", 18, GOLD)
    d.text(65, 665, 700, 80, "Department of Computer Science and Engineering\nKhulna University of Engineering and Technology", 17)
    d.text(65, 812, 715, 59, "C++17 + OpenGL 3.3\nProcedural models, mathematical motion, game-like controls", 15, CYAN)
    d.rect(965, 740, 515, 90, BG)
    d.text(990, 755, 475, 50, "36 crew   /   8 seats", 30, WHITE, True, DISPLAY)

    # 2. Required outline, with the game and engineering story established.
    d.new("Outline", "A world, its equations, the player's actions, and the frame that explains them.", "Presentation roadmap")
    agenda = [
        ("01", "Build the world", "Objects and screenshots; procedural curved surfaces", "Slides 3-5"),
        ("02", "Bring it to life", "Flight, machinery, workers, pursuit and attack", "Slide 6"),
        ("03", "Give the player control", "Keyboard + mouse, three cameras, winch and delivery", "Slide 7"),
        ("04", "Explain the result", "Rendering, feedback, audio and verification", "Slides 8-9"),
    ]
    for index, (number, heading, detail, page) in enumerate(agenda):
        y = 215 + index * 125
        d.text(65, y, 72, 55, number, 32, GOLD, True, DISPLAY)
        d.text(160, y + 2, 600, 42, heading, 24, WHITE, True)
        d.text(160, y + 47, 615, 44, detail, 17, MUTED)
        d.text(680, y + 5, 140, 30, page, 13, CYAN, align=PP_ALIGN.RIGHT)
        d.line(160, y + 104, 810, y + 104, EDGE, 0.8)
    d.text(925, 212, 570, 42, "THE RESCUE LOOP", 21, GOLD, True)
    steps = [("36 workers", "Leave the moving harvester"), ("8-seat cabin", "Hover and winch in small loads"), ("Cyan base", "Deliver: only then does score increase"), ("Threat + outcome", "Breach, extraction, debrief and retry")]
    for index, (heading, detail) in enumerate(steps):
        y = 278 + index * 105
        d.rect(925, y, 575, 80, PANEL)
        d.text(945, y + 8, 520, 40, heading, 23, CYAN, True)
        d.text(945, y + 49, 535, 28, detail, 16, WHITE)
    d.cue("SOURCE + AUDIO  >  GENERATED MESHES  >  INPUT / SIMULATION  >  CAMERA / AUDIO  >  GPU FRAME")

    # 3. First of exactly two object-inventory slides.
    d.new("Objects I / Desert and machines", "Every visible model is generated in code, assembled from indexed meshes, and shaded in real time.", "Objects / 1 of 2")
    objects = [
        ("01 / DUNE TERRAIN", "camera_0", "900 x 900 patch; 260 x 260 cells.\nHeight field, smooth normals, flat base terrace."),
        ("02 / ORNITHOPTER", "aircraft_crop", "Eight wings; hull/canopy, intakes, gimbals,\nturbines/nozzles, tail, sensor and fixed skids."),
        ("03 / SPICE HARVESTER", "harvester_crop", "Four crawlers, wheels/pads, cutter drum,\nexit/ramp, separators, bridge, stacks and lamps."),
        ("04 / THREE SPICE TANKS", "tank_crop", "Cradle/skids, cylinder vessel, domed ends,\nbands, hatch/valve and emissive indicator."),
        ("05 / ROCK FORMATIONS", "rocks_crop", "Eight sites x four pieces: irregular facets,\ndifferent scales, strata and procedural grain."),
        ("HOW THEY ARE ASSEMBLED", "aircraft_wire_crop", "Cube + cylinder + ellipsoid + wing mesh.\nParent/local transforms reuse GPU geometry."),
    ]
    for index, (heading, picture, detail) in enumerate(objects):
        column, row = index % 3, index // 3
        x, y = 60 + column * 505, 209 + row * 303
        d.text(x, y, 470, 30, heading, 17, GOLD, True)
        d.line(x, y + 34, x + 470, y + 34, EDGE, 0.7)
        d.image(picture, x, y + 45, 470, 178, fill=True)
        d.text(x, y + 231, 480, 63, detail, 15.5, WHITE)
    d.text(60, 825, 1460, 23, "Inspection close-ups reuse actual project draw routines. Terrain is a scene capture. No imported models or image textures.", 11, MUTED)

    # 4. Second object slide: modeled cavity, rescue infrastructure and interaction.
    d.new("Objects II / Rescue and sandworm", "The mission uses visible people, colour-coded markers, a winch line, and a genuinely recessed worm mouth.", "Objects / 2 of 2")
    d.image("mouth", 60, 208, 755, 424)
    d.text(80, 647, 725, 82, "A modeled cavity, not a black disc.\nBody + lip + funnel + deep throat + teeth", 23, WHITE, True)
    d.text(80, 733, 740, 48, "5 meshes  |  1,000 curved teeth  |  32 body ridge cycles", 17, GOLD)
    rescue = [
        ("WORKER", "worker_crop", "Torso, helmet/visor, limbs, locator"),
        ("SAFE BASE", "base_crop", "Pad, cyan ring/cross, four lights"),
        ("AMBER BEACON", "beacon_crop", "Slope-aligned ring, pole and lamp"),
        ("WINCH + LIFT", "winch_crop", "Cable + progress-based lift"),
    ]
    for index, (heading, picture, detail) in enumerate(rescue):
        x, y = 850 + (index % 2) * 345, 208 + (index // 2) * 280
        d.text(x, y, 325, 28, heading, 17, CYAN, True)
        d.image(picture, x, y + 37, 325, 180)
        d.text(x, y + 225, 326, 45, detail, 15, WHITE)
    d.cue("AMBER: awaiting rescue   |   CYAN: base / winch   |   WORM WAKE: four sand mounds + disturbed dust")

    # 5. Equations correspond to the actual surfaces, not decorative formulae.
    d.new("Curved geometry / The equations behind the objects", "Parameters control shape; triangles sample the equations; normals control lighting.", "Geometry and mathematical design")
    d.line(825, 202, 825, 780, EDGE, 0.8)
    d.line(60, 500, 1530, 500, EDGE, 0.8)
    d.text(60, 209, 735, 37, "DUNES / a multi-frequency height field", 20, GOLD, True)
    d.equation("terrain1", r"H_0(x,z)=28\sin(.009x+.004z)", 65, 260, 730, 47, 25)
    d.equation("terrain2", r"+19\cos(.004x-.012z)+5[1-|\sin(.028x+.016z)|]^2", 65, 313, 730, 45, 21)
    d.equation("terrain3", r"+1.8\sin(.065x+.038z)-4", 65, 365, 730, 44, 24)
    d.text(65, 425, 728, 64, "Broad dunes + sharper ridges + small undulations.\nBase flattening: smooth blend from radius 15 to 25.", 17, MUTED)
    d.text(865, 209, 655, 37, "ELLIPSOIDS / CYLINDERS / RINGS", 20, CYAN, True)
    d.equation("ellipsoid", r"\frac{x^2}{A^2}+\frac{y^2}{B^2}+\frac{z^2}{C^2}=1", 880, 258, 625, 63, 27)
    d.text(880, 322, 635, 25, "A/B/C: half-axis lengths; cylinder: unit height, radius 0.5.", 14, MUTED)
    d.equation("cylinder", r"x^2+z^2=.25,\quad -.5\leq y\leq .5", 880, 350, 625, 36, 24)
    d.equation("torus", r"P(a,b)=((1+.025\cos b)\cos a,\ .025\sin b,", 865, 394, 652, 35, 19)
    d.equation("torus2", r"(1+.025\cos b)\sin a)", 925, 436, 588, 36, 23)
    d.text(870, 473, 640, 24, "Torus a/b: major-circle / tube angles; radii 1 and 0.025.", 14, MUTED)
    d.text(60, 525, 730, 37, "WORM BODY / bend a radial surface", 20, GOLD, True)
    d.equation("body", r"P(t,a)=C(t)+\rho(t,a)(\cos a,-T_z\sin a,T_y\sin a)", 60, 581, 742, 55, 22)
    d.text(65, 655, 735, 131, "C(t): buried cubic Bezier + curved 64-unit neck.\nt: along body; a: ring angle; T: centreline tangent.\nrho: radius + 32 ridges/grain + mouth taper.\nProfile revolution creates lip, funnel and throat.", 17)
    d.text(865, 525, 650, 37, "TEETH / a curved, tapered centreline", 20, CYAN, True)
    d.equation("tooth1", r"P(t)=R(r-\ell t)+S(ct^2)", 870, 581, 630, 45, 26)
    d.equation("tooth2", r"+(0,d+.65\sin(\pi t)-Dt^2,0)", 870, 628, 630, 43, 24)
    d.equation("tooth3", r"q(t)=w(1-t)^{.85},\quad 0\leq t\leq 1", 870, 682, 630, 45, 25)
    d.text(865, 737, 650, 56, "R/S: radial/tangential; r/d: row radius/depth.\nl: reach; c: curl; D: drop; w: base width; t: base -> tip.", 14.5, MUTED)
    d.cue("A readable equation-to-object mapping: dunes, hull/tank/worker curves, beacon rings, flexible worm and curved teeth.")

    # 6. Motions shown by poses, equations and game-state transitions.
    d.new("Complex motion / One changing rescue scene", "Objects move independently, while mission events change what the player can rescue and what remains visible.", "Dynamic objects and dynamic scenes")
    stages = [("APPROACH", "worm_approach", "Wake + closing gap"), ("RISE", "worm_rise", "0-6 s after breach"), ("ATTACK / SWALLOW", "worm_attack", "6-12 s: lunge + sink"), ("RETREAT", "worm_retreat", "13.5-18 s: withdraw")]
    for index, (title, picture, detail) in enumerate(stages):
        x = 60 + index * 375
        d.text(x, 207, 350, 31, title, 19, GOLD if index < 2 else CYAN, True)
        d.image(picture, x, 246, 350, 185)
        d.text(x, 442, 350, 35, detail, 16, MUTED)
        if index < 3: d.arrow(x + 351, 322, 19)
    d.line(60, 489, 1530, 489, EDGE)
    d.text(60, 509, 705, 32, "FLIGHT + ARTICULATION", 20, CYAN, True)
    d.equation("velocity", r"v_{\mathrm{next}}=v+(v_d-v)(1-e^{-k\Delta t})", 62, 553, 708, 45, 24)
    d.equation("wing", r"\beta_1=34\sin\phi,\quad \beta_2=34\sin(\phi+.75\pi)", 62, 607, 708, 45, 23)
    d.text(65, 664, 715, 123, "k = 5.5 normal / 12 braking; speeds 28 / 48 units/s.\nTerrain following + pitch/bank; analytically smoothed yaw.\nWing angles in degrees; phase 23 -> 36 rad/s; twist/exhaust.\nTracks/drum follow distance travelled, not idle clock time.", 16)
    d.text(840, 509, 690, 32, "AUTONOMOUS ROUTE + EVACUATION", 20, GOLD, True)
    d.equation("route1", r"x=-40+.62t+15\sin(.045t)", 840, 553, 682, 40, 24)
    d.equation("route2", r"z=-55+22\sin(.032t)+9\sin(.071t)", 840, 602, 682, 40, 24)
    d.text(845, 657, 681, 126, "t: seconds; tangent heading and slope alignment.\nGroups release at 6 + 7j s; crew run at 3.8 units/s.\nWorm gap: 220 -> 18; limbs animate; radius/height hazards.\nCurrent breach: 150 s; then 60 s final extraction.", 16.5)
    d.cue("TITLE  >  FLIGHT / EVACUATION  >  BREACH / HARVESTER LOSS  >  EXTRACTION  >  DEBRIEF / RETRY / NEXT")

    # 7. All user bindings, three cameras, and measurable interaction limits.
    d.new("Game-like controls / Fly, inspect, rescue", "The aircraft and the mission respond to input; changing view does not change the aircraft's heading.", "Controllability and camera movement")
    d.text(60, 207, 700, 35, "FLIGHT INPUT", 21, GOLD, True)
    controls = [("W / S", "Forward / reverse thrust"), ("A / D", "Strafe without turning the nose"), ("Q / E", "Raise / lower terrain-following altitude"), ("MOUSE", "Hover steers yaw + camera pitch; centre dead zone"), ("ARROWS", "Left/right turn; up/down camera elevation"), ("SHIFT", "Rechargeable boost; release after depletion"), ("SPACE", "Brake + descend + winch / unload")]
    for index, (key, action) in enumerate(controls):
        y = 263 + index * 58
        d.rect(60, y, 132, 42, PANEL)
        d.text(70, y + 8, 114, 27, key, 16, CYAN, True)
        d.text(210, y + 8, 590, 43, action, 17)
    d.text(60, 701, 737, 88, "Safety: altitude 4-55; rock/harvester clearance help.\nFocus loss pauses and clears held input; reset/resize\nrecentres the cursor. Movement stays inside +/-160.", 16, MUTED)
    d.text(845, 207, 680, 35, "C / THREE HEADING-FOLLOWING VIEWS", 20, CYAN, True)
    for mode, label in enumerate(["WIDE / 40", "CLOSE / 20", "TACTICAL / 100 UP"]):
        x = 845 + mode * 234
        d.image(f"camera_{mode}", x, 260, 216, 139)
        d.text(x, 410, 220, 32, label, 13, GOLD, True)
    d.text(845, 454, 680, 54, "Smooth follow + terrain-safe eye; 58-degree perspective.\nAll are third-person, including the close preset.", 16, MUTED)
    d.line(845, 521, 1530, 521, EDGE)
    d.text(845, 540, 681, 38, "SPACE / RESCUE IS A STATE CHANGE", 20, WHITE, True)
    d.text(845, 591, 680, 100, "Pickup: distance <12, height <=11, speed <6.\nHold 0.42 s/person; cabin capacity = 8.\nDeliver: base distance <14; hold 1.2 s per load.", 18)
    d.text(845, 704, 680, 61, "Inside -> Running/Waiting -> Aboard -> Safe\nOnly Safe increases score; danger/timeout can cause Lost.", 16, CYAN)
    d.cue("ENTER start/next   |   P pause   |   R retry   |   M mute   |   F wireframe   |   F11 fullscreen   |   ESC exit")

    # 8. Every rendering extra tied to visible evidence and a GPU stage.
    d.new("Rendering / What makes the scene convincing", "A forward OpenGL renderer: one shadow target, then the visible frame, particles, and the interface.", "Additional graphics features")
    pipeline = [("01 / SHADOW DEPTH", "2048 x 2048 depth map"), ("02 / SKY + SUN", "Generated sky view rays"), ("03 / LIT OBJECTS", "Materials + shadows + fog"), ("04 / DUST", "Sorted alpha billboards"), ("05 / HUD", "Screen-space triangles")]
    for index, (title, detail) in enumerate(pipeline):
        x = 60 + index * 300
        d.rect(x, 213, 276, 95, PANEL)
        d.text(x + 13, 225, 253, 34, title, 16, CYAN if index % 2 else GOLD, True)
        d.text(x + 13, 264, 255, 34, detail, 14, WHITE)
        if index < 4: d.arrow(x + 277, 245, 19)
    columns = [
        ("SHADOWS + MATERIALS", "shadow_detail_crop", "Directional Blinn-Phong + sand-wrap diffuse.\n5 x 5 PCF; receiver-plane correction + bias.\nGlossy opaque surfaces; emissive ornaments."),
        ("PROCEDURAL SURFACE DETAIL", "rocks_crop", "Sand ripples, bump-like normals and sparkle.\nRock strata/grain; no bitmap texture files.\nSky gradient, sun disc/halo, distance fog, tone mapping."),
        ("DUST + COLLAPSING TERRAIN", "dust_detail_crop", "2,400 particles: wind/crawler/worm/downwash.\nSorted instances + lifetime/alpha fades.\nDepth-tested, camera-facing billboards.\nCrater: 13 units deep, tapers to radius 48."),
    ]
    for index, (heading, picture, detail) in enumerate(columns):
        x = 60 + index * 505
        d.text(x, 342, 480, 34, heading, 18, GOLD, True)
        if index == 2:
            d.image(picture, x, 391, 232, 211)
            d.image("crater", x + 242, 391, 232, 211)
            d.rect(x, 573, 232, 29, BG)
            d.rect(x + 242, 573, 232, 29, BG)
            d.text(x + 8, 579, 211, 24, "BREACH DUST", 12, WHITE, True)
            d.text(x + 250, 579, 215, 24, "GPU DEPRESSION", 12, WHITE, True)
        else:
            d.image(picture, x, 391, 474, 211)
        d.text(x, 620, 480, 125, detail, 16.5)
    d.text(63, 758, 1470, 32, "Depth/culling + requested 4x MSAA; inverse-transpose normals. Crater is GPU-only: CPU ground/collision height is unchanged.", 15, MUTED)
    d.cue("DEMO: inspect long shadows + surface detail; press F for triangle topology; show breach dust and the visual depression.")

    # 9. Feedback, eight sounds, complete engineering and honest current test results.
    d.new("Feedback / The player can read and hear the mission", "Custom HUD, spatial sound, runtime safeguards, and reproducible build/testing tools.", "Additional interface, audio and engineering features")
    d.text(60, 209, 845, 37, "INSTRUMENTATION / 5 x 7 BITMAP FONT", 20, CYAN, True)
    d.image("status_crop", 60, 259, 845, 86, False)
    d.text(60, 358, 850, 41, "Crew / cabin / score / wave; breach or extraction timer; seat and progress bars.", 16)
    d.image("telemetry_crop", 60, 416, 300, 171)
    d.image("radar_crop", 610, 417, 236, 294)
    d.text(65, 603, 496, 120, "Altitude, speed, FPS, boost, threat gap.\nNorth-up radar: cyan aircraft/base; amber crew;\nwhite harvester; red worm; world range +/-300.\nAdaptive layout and camera guide brackets.", 16)
    d.image("action_crop", 60, 733, 845, 69)
    d.text(60, 809, 850, 34, "Context prompts; title/pause/debrief panels; delivery-only score; session best.", 14, MUTED)
    d.line(944, 209, 944, 827, EDGE)
    d.text(980, 210, 555, 35, "8 SPATIAL SOUND EFFECTS", 21, GOLD, True)
    d.text(982, 264, 553, 118, "Loops: flight / boost / harvester / worm rumble.\nEvents: breach / winch / safe delivery / crew help.\nCamera listener; distance attenuation; speed pitch;\nboost/threat volume; mute and pause silencing.", 16.5)
    d.text(982, 387, 555, 67, "Silent fallback; focus pause and cursor recentering.\nFullscreen/resize; minimized waiting; substeps.", 16, MUTED)
    d.line(980, 469, 1530, 469, EDGE)
    d.text(980, 488, 555, 36, "ENGINEERING + CURRENT VERIFICATION", 19, CYAN, True)
    d.text(980, 542, 555, 120, "C++17 / GLFW / GLAD / GLM / miniaudio.\nCMake + Visual Studio; audio copied post-build.\nProcedural startup meshes and embedded GLSL.\nReport PDF + source documentation/manifest tooling.", 16.5)
    d.text(980, 669, 550, 41, "4 / 4 graphical smoke modes passed", 21, WHITE, True)
    d.text(980, 721, 550, 42, "23 / 27 mission cases; 1,057 checks", 19, GOLD, True)
    d.text(980, 774, 554, 63, "Current timer is 150 s; five assertions fail because\ntests expect 210-to-150 progression. Source untouched.", 15, RED)

    # 10. Required final Thank You slide, closing the exact ten-slide sequence.
    d.new("Thank You", chrome=False)
    d.image("pursuit", 0, 0, 1600, 900)
    d.rect(0, 0, 1600, 270, BG)
    d.rect(0, 565, 1600, 335, BG)
    d.text(65, 44, 1480, 31, "10 / CLOSING REMARKS", 13, GOLD, True)
    d.text(60, 109, 1460, 126, "THANK YOU", 74, WHITE, True, DISPLAY)
    d.text(65, 590, 1470, 55, "Procedural geometry. Mathematical motion. A controllable rescue mission.", 26, WHITE, True)
    d.text(65, 668, 1440, 55, "Asique Ehetasamul Haque  |  Roll 2107096  |  CSE - 4102  |  Year 4th, Semester 1st, Group B2", 18, GOLD)
    d.text(65, 745, 1420, 45, "Objects, curved equations, dynamic scenes, cameras, interaction, shadows, surface detail, dust, HUD and sound.", 18, WHITE)
    d.text(65, 811, 1450, 52, "Project report available. Your completed demonstration video is separate; no video is embedded in this deck.\nLimits: arcade physics; opaque glass; GPU-only crater. 'Karpov objects' is undefined and is not claimed.", 14, MUTED)

    if len(d.slides) != 10: raise RuntimeError("Exactly ten slides are required.")
    for number, slide in enumerate(d.slides, 1):
        slide.notes_slide.notes_text_frame.text = NOTES[number]
    d.presentation.save(HERE / "Arrakis_Presentation.pptx")
    hashes = {str(path.relative_to(ROOT)).replace("\\", "/"): hashlib.sha256(path.read_bytes()).hexdigest()
              for path in [ROOT / "src/arrakis.cpp", ROOT / "src/arrakis_game.h", ROOT / "src/arrakis_worm.h", ROOT / "src/arrakis_hud.h", ROOT / "src/arrakis_audio.h"]}
    manifest = {"slide_count": 10, "aspect_ratio": "16:9", "source_sha256": hashes,
                "video_embedded": False, "slides": d.manifest, "feature_coverage": FEATURE_COVERAGE,
                "capture_provenance": "Gameplay: current Release smoke modes. Close-ups/poses: capture.cpp reuses original draw routines with staged state and inspection framing; CPU/GPU project source is not edited.",
                "verification": {"mission_cases_passed": 23, "mission_cases_total": 27, "checks": 1057, "failed_assertions": 5, "smoke_modes_passed": 4, "frames_per_mode": 120}}
    (HERE / "presentation_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    write_presenter_guide(d.manifest)
    print("Generated ten-slide editable PowerPoint, presenter notes PDF and feature/source manifest.")


def write_presenter_guide(slides):
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="GuideBody", fontName="Helvetica", fontSize=10.3, leading=14.5, spaceAfter=8, alignment=TA_LEFT))
    styles.add(ParagraphStyle(name="GuideHeader", fontName="Helvetica-Bold", fontSize=19, leading=23, textColor=colors.HexColor("#153D45"), spaceAfter=12))
    styles.add(ParagraphStyle(name="GuideSub", fontName="Helvetica-Bold", fontSize=11, leading=15, spaceBefore=6, spaceAfter=4))
    document = SimpleDocTemplate(str(HERE / "Presenter_Notes.pdf"), pagesize=(595.28, 841.89), leftMargin=45, rightMargin=45, topMargin=42, bottomMargin=42,
                                 title="Arrakis - Presenter Notes", author="Asique Ehetasamul Haque")
    story = []
    for slide in slides:
        number = slide["number"]
        story.append(Paragraph(f"Slide {number:02d} / {escape(slide['title'])}", styles["GuideHeader"]))
        for paragraph in NOTES[number].split("\n\n"):
            value = escape(paragraph).replace("\n", "<br/>")
            story.append(Paragraph(value, styles["GuideSub"] if len(paragraph) < 70 and paragraph.endswith(":") else styles["GuideBody"]))
        if number != 10: story.append(PageBreak())
    document.build(story)


if __name__ == "__main__":
    create_deck()
